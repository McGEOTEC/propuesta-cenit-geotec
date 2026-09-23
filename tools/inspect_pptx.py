import sys
from pptx import Presentation
import json

prs_path = r'F:\PROYECTOS\EMPRESARIAL\GEOTEC_ESTRATEGIA_COMERCIAL\ESTILO_GEOTEC\GEOTEC_PLANTILLA_BASE_PRESENTACIONES.pptx'
prs = Presentation(prs_path)

print(f"Total slides in template: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
    print(f"\n--- Slide {i+1} ---")
    print(" | ".join(texts[:5]))
